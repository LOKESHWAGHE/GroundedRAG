import pytest
from docx import Document as DocxDocument
from reportlab.pdfgen import canvas

from app.services.ingestion.parsers import ParseError, parse_document


def make_pdf(path, page_texts):
    c = canvas.Canvas(str(path))
    for text in page_texts:
        if text:
            c.drawString(72, 720, text)
        c.showPage()
    c.save()


def test_txt_is_single_page(tmp_path):
    p = tmp_path / "a.txt"
    p.write_text("hello world", encoding="utf-8")
    pages = parse_document(p, ".txt")
    assert [(x.page_number, x.text) for x in pages] == [(1, "hello world")]


def test_markdown_keeps_headings(tmp_path):
    p = tmp_path / "a.md"
    p.write_text("# Title\n\nBody", encoding="utf-8")
    assert parse_document(p, ".md")[0].text == "# Title\n\nBody"


def test_invalid_utf8_raises(tmp_path):
    p = tmp_path / "a.txt"
    p.write_bytes(b"\xff\xfe\xfa bad")
    with pytest.raises(ParseError):
        parse_document(p, ".txt")


def test_pdf_pages_are_numbered_from_one(tmp_path):
    p = tmp_path / "a.pdf"
    make_pdf(p, ["First page text", "Second page text"])
    pages = parse_document(p, ".pdf")
    assert [x.page_number for x in pages] == [1, 2]
    assert "First page text" in pages[0].text
    assert "Second page text" in pages[1].text


def test_pdf_blank_page_keeps_its_page_number(tmp_path):
    p = tmp_path / "a.pdf"
    make_pdf(p, ["one", "", "three"])
    pages = parse_document(p, ".pdf")
    assert len(pages) == 3 and pages[1].text.strip() == ""


def test_corrupt_pdf_raises_parse_error(tmp_path):
    p = tmp_path / "a.pdf"
    p.write_bytes(b"%PDF-1.4 this is garbage")
    with pytest.raises(ParseError):
        parse_document(p, ".pdf")


def test_docx_headings_and_tables(tmp_path):
    p = tmp_path / "a.docx"
    d = DocxDocument()
    d.add_heading("Overview", level=1)
    d.add_paragraph("Intro paragraph.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Product"
    t.rows[0].cells[1].text = "Database"
    d.add_heading("Details", level=2)
    d.save(str(p))
    text = parse_document(p, ".docx")[0].text
    assert "# Overview" in text
    assert "Intro paragraph." in text
    assert "Product | Database" in text
    assert "## Details" in text


def test_corrupt_docx_raises_parse_error(tmp_path):
    p = tmp_path / "a.docx"
    p.write_bytes(b"PK\x03\x04 not really a docx")
    with pytest.raises(ParseError):
        parse_document(p, ".docx")
